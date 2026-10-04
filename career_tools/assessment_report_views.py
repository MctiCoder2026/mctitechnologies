from datetime import date, datetime, time, timedelta
from io import BytesIO
from xml.sax.saxutils import escape

from django.contrib.auth.decorators import login_required
from django.core.exceptions import ObjectDoesNotExist
from django.core.paginator import Paginator
from django.db.models import CharField, Value
from django.db.models.functions import Coalesce, Lower, NullIf, Trim
from django.http import HttpResponse, HttpResponseBadRequest, HttpResponseForbidden
from django.shortcuts import render
from django.urls import reverse
from django.utils import timezone

from .models import AptitudeAttempt, CareerCompassAttempt
from assessments.models import AssessmentAttempt, HiringApplication

BRANCHES = (
    "kharghar", "panvel", "koperkhairane", "kamothe",
    "ghansoli", "nerul", "head_office",
)
KINDS = {
    "aptitude": "Aptitude",
    "compass": "Career Compass",
    "ai": "AI Ready",
    "scholarship": "Scholarship",
    "cyber": "Cyber Awareness",
    "hiring": "Hiring",
}


def _branch_expression(*fields):
    return Lower(Trim(Coalesce(
        *[NullIf(Trim(field), Value("")) for field in fields],
        Value("unassigned"), output_field=CharField(),
    )))


def _related(obj, name):
    try:
        return getattr(obj, name)
    except ObjectDoesNotExist:
        return None


def _dates(params):
    today = timezone.localdate()
    period = params.get("period", "today")
    if period == "today":
        return today, today, period
    if period == "week":
        return today - timedelta(days=6), today, period
    if period == "all":
        return None, None, period
    if period != "custom":
        raise ValueError("Invalid date filter.")
    start = date.fromisoformat(params.get("start", ""))
    end = date.fromisoformat(params.get("end", ""))
    if start > end:
        raise ValueError("Start date must be before end date.")
    return start, end, period


def _filter(qs, date_field, start, end, branch):
    if branch:
        qs = qs.filter(report_branch=branch)
    if start:
        zone = timezone.get_current_timezone()
        lower = timezone.make_aware(datetime.combine(start, time.min), zone)
        upper = timezone.make_aware(
            datetime.combine(end + timedelta(days=1), time.min), zone
        )
        qs = qs.filter(**{
            date_field + "__gte": lower,
            date_field + "__lt": upper,
        })
    return qs


def _row(kind, obj, name, mobile, when, score, detail, url=""):
    return {
        "kind": kind, "label": KINDS[kind], "name": name,
        "mobile": mobile, "branch": obj.report_branch,
        "date": when, "score": score, "detail": detail, "url": url,
        "id": obj.pk,
    }


def _collect(start, end, branch, selected_kind, management, search):
    rows = []
    if selected_kind in ("", "aptitude"):
        qs = AptitudeAttempt.objects.filter(status="completed").select_related(
            "profile", "profile__enquiry", "profile__student"
        ).annotate(report_branch=_branch_expression(
            "profile__preferred_branch", "profile__enquiry__branch",
            "profile__student__branch",
        ))
        for a in _filter(qs, "completed_at", start, end, branch).iterator():
            detail = (
                f"Correct: {a.correct_answers}/{a.total_questions}; "
                f"Numerical: {a.numerical_score}%; Logical: {a.logical_score}%; "
                f"Verbal: {a.verbal_score}%; Computer: {a.computer_score}%; "
                f"Career: {a.career_score}%"
            )
            rows.append(_row(
                "aptitude", a, a.profile.full_name, a.profile.mobile,
                a.completed_at, a.score_percentage, detail,
                reverse("career_tools:counsellor_profile_detail",
                        args=[a.profile_id]),
            ))

    if selected_kind in ("", "compass"):
        qs = CareerCompassAttempt.objects.filter(
            status="completed"
        ).select_related("profile", "profile__enquiry", "profile__student").annotate(
            report_branch=_branch_expression(
                "profile__preferred_branch", "profile__enquiry__branch",
                "profile__student__branch",
            )
        )
        for a in _filter(qs, "completed_at", start, end, branch).iterator():
            detail = (
                f"Directions: {a.top_direction}, {a.second_direction}, "
                f"{a.third_direction}; Counselling: "
                f"{a.get_counselling_status_display()}"
            )
            rows.append(_row(
                "compass", a, a.profile.full_name, a.profile.mobile,
                a.completed_at, None, detail,
                reverse("career_tools:counsellor_profile_detail",
                        args=[a.profile_id]),
            ))

    slugs = {
        "ai": [
            "ai-ready-maharashtra-2026",
            "ai-ready-maharashtra-2026-marathi",
        ],
        "cyber": ["mcti-cyber-fraud-awareness-2026"],
    }
    for kind in ("ai", "scholarship", "cyber"):
        if selected_kind not in ("", kind):
            continue
        qs = AssessmentAttempt.objects.filter(status="completed").select_related(
            "assessment", "participant_profile",
            "participant_profile__enquiry", "result", "scholarship_result",
        )
        if kind == "scholarship":
            qs = qs.filter(scholarship_result__isnull=False)
        else:
            qs = qs.filter(assessment__slug__in=slugs[kind])
        qs = qs.annotate(report_branch=_branch_expression(
            "scholarship_result__preferred_branch",
            "participant_profile__enquiry__branch",
        ))
        for a in _filter(qs, "completed_at", start, end, branch).iterator():
            result = _related(a, "result")
            detail = f"Marks: {a.score}/{a.total_marks}"
            if result:
                detail += f"; Result: {result.result_title or result.result_band}"
            scholarship = _related(a, "scholarship_result")
            if kind == "scholarship" and scholarship:
                detail += (
                    f"; Scholarship: {scholarship.scholarship_percentage}%; "
                    f"Course: {scholarship.selected_course}; "
                    f"Status: {scholarship.get_claim_status_display()}"
                )
            rows.append(_row(
                kind, a, a.participant_name, a.participant_mobile,
                a.completed_at, a.percentage, detail,
            ))

    if management and selected_kind in ("", "hiring") and branch in ("", "unassigned"):
        qs = HiringApplication.objects.select_related(
            "candidate", "job_role", "assessment_attempt"
        ).annotate(report_branch=Value("unassigned", output_field=CharField()))
        for a in _filter(qs, "created_at", start, end, branch).iterator():
            attempt = a.assessment_attempt
            score = attempt.percentage if attempt and attempt.status == "completed" else None
            detail = (
                f"Role: {a.job_role}; Status: {a.get_status_display()}; "
                f"Employment: {a.get_employment_type_display()}"
            )
            rows.append(_row(
                "hiring", a, a.candidate.full_name, a.candidate.mobile,
                a.created_at, score, detail,
            ))

    if search:
        needle = search.casefold()
        rows = [r for r in rows if needle in (
            r["name"] + " " + r["mobile"] + " " + r["detail"]
        ).casefold()]
    rows.sort(key=lambda r: (
        r["date"].timestamp() if r["date"] else 0, r["id"]
    ), reverse=True)
    return rows


def _export(rows, format_name, title):
    headers = ["Date", "Assessment", "Name", "Mobile", "Branch", "Score (%)", "Performance"]
    values = [
        [
            timezone.localtime(r["date"]).strftime("%d %b %Y %H:%M") if r["date"] else "",
            r["label"], r["name"], r["mobile"], r["branch"],
            float(r["score"]) if r["score"] is not None else "",
            r["detail"],
        ]
        for r in rows
    ]
    if format_name == "xlsx":
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        wb = Workbook()
        ws = wb.active
        ws.title = "Assessment Reports"
        ws.append(headers)
        for data in values:
            ws.append(data)
            for cell in ws[ws.max_row]:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="171717")
        for column, width in zip("ABCDEFG", (22, 22, 26, 18, 20, 15, 90)):
            ws.column_dimensions[column].width = width
        for row in ws.iter_rows(min_row=2):
            row[-1].alignment = Alignment(wrap_text=True, vertical="top")
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        output = BytesIO()
        wb.save(output)
        response = HttpResponse(
            output.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
    else:
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, LongTable, TableStyle
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib import colors
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        font_path = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
        font = "Helvetica"
        if __import__("os").path.exists(font_path):
            if "MCTIReport" not in pdfmetrics.getRegisteredFontNames():
                pdfmetrics.registerFont(TTFont("MCTIReport", font_path))
            font = "MCTIReport"
        styles = getSampleStyleSheet()
        small = ParagraphStyle("ReportCell", fontName=font, fontSize=7, leading=10)
        output = BytesIO()
        doc = SimpleDocTemplate(
            output, pagesize=landscape(A4),
            rightMargin=22, leftMargin=22, topMargin=25, bottomMargin=30,
        )
        story = [
            Paragraph("MCTI Technologies — Assessment Reports", styles["Title"]),
            Paragraph(escape(title), styles["Normal"]),
            Paragraph(f"Records: {len(rows)}", styles["Normal"]),
            Spacer(1, 12),
        ]
        data = [[Paragraph(escape(str(v)), small) for v in headers]]
        data += [[Paragraph(escape(str(v)), small) for v in row] for row in values]
        table = LongTable(data, repeatRows=1, colWidths=[80, 80, 100, 75, 75, 55, 330])
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#ff7200")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#dddddd")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f7f7f7")]),
            ("TOPPADDING", (0, 0), (-1, -1), 6),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
        ]))
        # Fit the table inside landscape A4.
        table._argW = [w * (landscape(A4)[0] - 44) / 795 for w in table._argW]
        story.append(table)
        def footer(canvas, document):
            canvas.setFont("Helvetica", 8)
            canvas.drawString(22, 15, "An Initiative by Maharashtra Computer Training Institute")
            canvas.drawRightString(landscape(A4)[0]-22, 15, f"Page {document.page}")
        doc.build(story, onFirstPage=footer, onLaterPages=footer)
        response = HttpResponse(output.getvalue(), content_type="application/pdf")
    response["Content-Disposition"] = (
        f'attachment; filename="MCTI-Assessment-Reports.{format_name}"'
    )
    response["Cache-Control"] = "private, no-store"
    return response


@login_required(login_url="staff_login")
def assessment_reports(request):
    user = request.user
    if not user.is_active or not (user.is_staff or user.is_superuser):
        return HttpResponseForbidden("Staff access required.")
    staff = _related(user, "staff_profile")
    management = user.is_superuser or bool(
        staff and staff.is_active
        and (staff.branch or "").strip().lower() == "head_office"
    )
    if not management and not (staff and staff.is_active and staff.branch):
        return HttpResponseForbidden("An active staff branch is required.")
    branch = (
        request.GET.get("branch", "").strip().lower()
        if management else staff.branch.strip().lower()
    )
    if branch and branch not in (*BRANCHES, "unassigned"):
        return HttpResponseBadRequest("Invalid branch.")
    kind = request.GET.get("kind", "").strip()
    if kind and (kind not in KINDS or (kind == "hiring" and not management)):
        return HttpResponseForbidden("This report is not available.")
    try:
        start, end, period = _dates(request.GET)
    except ValueError:
        return HttpResponseBadRequest("Choose valid start and end dates.")
    search = request.GET.get("q", "").strip()
    rows = _collect(start, end, branch, kind, management, search)
    title = (
        f"Branch: {branch.title() if branch else 'All branches'} | "
        f"Period: {start or 'All dates'} to {end or 'All dates'}"
    )
    export = request.GET.get("export", "")
    if export:
        if export not in ("xlsx", "pdf"):
            return HttpResponseBadRequest("Invalid export format.")
        return _export(rows, export, title)
    params = request.GET.copy()
    params.pop("page", None)
    params.pop("export", None)
    summary = {}
    for row in rows:
        summary[row["branch"]] = summary.get(row["branch"], 0) + 1
    cards = [
        {"key": key, "label": label, "count": sum(r["kind"] == key for r in rows)}
        for key, label in KINDS.items() if management or key != "hiring"
    ]
    response = render(request, "career_tools/assessment_reports.html", {
        "page": Paginator(rows, 50).get_page(request.GET.get("page")),
        "total": len(rows), "cards": cards,
        "summary": sorted(summary.items()),
        "management": management, "branches": (*BRANCHES, "unassigned"),
        "selected_branch": branch, "selected_kind": kind,
        "period": period, "start": start, "end": end, "search": search,
        "params": params.urlencode(), "title": title,
    })
    response["Cache-Control"] = "private, no-store"
    return response
