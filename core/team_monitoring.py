from datetime import date, datetime, time, timedelta
from decimal import Decimal
from io import BytesIO
from xml.sax.saxutils import escape

from django.contrib import messages
from django.contrib.auth import get_user_model
from django.contrib.auth.decorators import login_required
from django.db import transaction
from django.db.models import Q
from django.http import HttpResponse, HttpResponseForbidden
from django.shortcuts import render, redirect
from django.utils import timezone

from .models import (
    Enquiry, EnquiryActivity, Admission, FeePayment,
    StaffLoginLog, StaffPerformanceReward, StaffAttendance,
)

@login_required(login_url="staff_login")
def team_monitoring(request):
    if not request.user.is_active or not request.user.is_superuser:
        return HttpResponseForbidden("Power User access required.")

    from .views import record_admin_attendance
    record_admin_attendance(request, source="power_user_session")

    today = timezone.localdate()
    period = request.GET.get("period", "week")
    start, end = today - timedelta(days=6), today
    try:
        if period == "today":
            start = end = today
        elif period == "month":
            start = today.replace(day=1)
        elif period == "custom":
            start = date.fromisoformat(request.GET.get("start", ""))
            end = date.fromisoformat(request.GET.get("end", ""))
        else:
            period = "week"
        if start > end or (end-start).days > 92 or end > today:
            raise ValueError()
    except ValueError:
        return HttpResponse("Choose a valid date range of up to 93 days.", status=400)

    zone = timezone.get_current_timezone()
    lower = timezone.make_aware(datetime.combine(start, time.min), zone)
    upper = timezone.make_aware(datetime.combine(end+timedelta(days=1), time.min), zone)

    users = list(get_user_model().objects.filter(
        Q(is_staff=True) | Q(is_superuser=True)
    ).select_related("staff_profile").order_by("username"))
    branch = request.GET.get("branch", "").strip().lower()
    from .student_attendance_awards import attendance_awards
    award_context = attendance_awards(request, branch)
    if isinstance(award_context, HttpResponse):
        return award_context

    branch_choices = sorted({
        (getattr(getattr(u, "staff_profile", None), "branch", "") or "").strip().lower()
        for u in users
    } - {""})
    if branch:
        users = [
            u for u in users
            if (getattr(getattr(u, "staff_profile", None), "branch", "") or "").strip().lower() == branch
        ]
    ids = [u.pk for u in users]
    stats = {
        uid: {
            "calls": set(), "whatsapp": set(), "contacts": set(),
            "results": set(), "conversions": set(),
            "days": set(), "latest": None,
            "admissions": 0, "receipts": 0,
            "collection": Decimal("0"), "logins": 0, "denied": 0,
            "login_days": set(), "today": 0, "overdue": 0,
        } for uid in ids
    }

    def record_work(uid, stamp):
        item = stats[uid]
        item["days"].add(timezone.localtime(stamp).date())
        if item["latest"] is None or stamp > item["latest"]:
            item["latest"] = stamp

    activities = EnquiryActivity.objects.filter(
        created_by_id__in=ids, created_at__gte=lower, created_at__lt=upper,
    ).exclude(
        activity_type="note", message="New enquiry received from website."
    )
    for a in activities.values(
        "created_by_id", "enquiry_id", "activity_type",
        "campaign_key", "message", "created_at",
    ):
        uid = a["created_by_id"]
        item = stats[uid]
        record_work(uid, a["created_at"])
        kind = a["activity_type"]
        if kind in ("call", "whatsapp"):
            item["calls" if kind == "call" else "whatsapp"].add(a["enquiry_id"])
        if kind == "converted":
            item["conversions"].add(a["enquiry_id"])
        if kind == "followup" and a["campaign_key"] == "followup_completed":
            item["results"].add(a["enquiry_id"])
            if (a["message"] or "").startswith((
                "Follow-up result recorded: Contacted successfully.",
                "Follow-up result recorded: WhatsApp message sent",
            )):
                item["contacts"].add(a["enquiry_id"])

    for a in Admission.objects.filter(
        created_by_id__in=ids, created_at__gte=lower, created_at__lt=upper,
    ).values("created_by_id", "created_at"):
        uid = a["created_by_id"]
        stats[uid]["admissions"] += 1
        record_work(uid, a["created_at"])

    for p in FeePayment.objects.filter(
        collected_by_id__in=ids, created_at__gte=lower, created_at__lt=upper,
    ).values("collected_by_id", "amount", "created_at"):
        uid = p["collected_by_id"]
        stats[uid]["receipts"] += 1
        stats[uid]["collection"] += p["amount"]
        record_work(uid, p["created_at"])

    for log in StaffLoginLog.objects.filter(
        user_id__in=ids, created_at__gte=lower, created_at__lt=upper,
    ).values("user_id", "status", "created_at"):
        item = stats[log["user_id"]]
        if log["status"] == "success":
            item["logins"] += 1
            item["login_days"].add(timezone.localtime(log["created_at"]).date())
        else:
            item["denied"] += 1

    pending = Enquiry.objects.exclude(status__in=("converted", "closed"))
    if branch:
        pending = pending.filter(branch__iexact=branch)
    for e in pending.filter(assigned_user_id__in=ids).values(
        "assigned_user_id", "followup_date",
    ):
        due = e["followup_date"]
        if due == today:
            stats[e["assigned_user_id"]]["today"] += 1
        elif due and due < today:
            stats[e["assigned_user_id"]]["overdue"] += 1


    # Integrated existing staff attendance
    attendance_map = {uid: {} for uid in ids}
    attendance_records = StaffAttendance.objects.filter(
        staff__user_id__in=ids,
    ).filter(
        Q(attendance_date__gte=start, attendance_date__lte=end) |
        Q(attendance_date=today)
    ).select_related("staff", "staff__user")
    for entry in attendance_records:
        attendance_map[entry.staff.user_id][entry.attendance_date] = entry

    reward_map = {
        r.user_id: r for r in StaffPerformanceReward.objects.filter(
            user_id__in=ids, period_start=start, period_end=end,
        )
    }
    rows = []
    for u in users:
        item = stats[u.pk]
        profile = getattr(u, "staff_profile", None)
        joined = timezone.localtime(u.date_joined).date()
        observable_start = max(start, joined)
        observable_days = max(0, (end-observable_start).days+1)
        work_days = {d for d in item["days"] if observable_start <= d <= end}

        attendance = attendance_map[u.pk]
        period_records = [
            a for day, a in attendance.items() if start <= day <= end
        ]
        attendance_today = attendance.get(today)
        attendance_start = max(
            start, profile.joining_date if profile else joined, joined
        )
        possible_days = max(0, (end-attendance_start).days+1)
        recorded_days = sum(
            attendance_start <= a.attendance_date <= end for a in period_records
        )
        score = len(item["contacts"])*2 + item["admissions"]*10
        row = {
            "user": u, "name": u.get_full_name() or u.username,
            "branch": getattr(profile, "branch", "") or "—",
            "role": "Superuser" if u.is_superuser else (
                getattr(profile, "designation", "") or "Staff"
            ),
            "enabled": u.is_active and (
                u.is_superuser or bool(profile and profile.is_active)
            ),
            "calls": len(item["calls"]), "whatsapp": len(item["whatsapp"]),
            "contacts": len(item["contacts"]), "results": len(item["results"]),
            "conversions": len(item["conversions"]),
            "admissions": item["admissions"],
            "receipts": item["receipts"], "collection": item["collection"],
            "days": len(work_days),
            "quiet_days": max(0, observable_days-len(work_days)),
            "latest": item["latest"], "logins": item["logins"],
            "login_days": len(item["login_days"]), "denied": item["denied"],
            "today": item["today"], "overdue": item["overdue"],
            "attendance_today": attendance_today,
            "attendance_status": (
                attendance_today.get_status_display() if attendance_today
                else ("Not recorded" if profile else "No staff profile")
            ),
            "present_days": sum(a.status == "present" for a in period_records),
            "leave_days": sum(a.status == "leave" for a in period_records),
            "absent_days": sum(a.status == "absent" for a in period_records),
            "missing_days": max(0, possible_days-recorded_days) if profile else None,
            "score": score, "reward": reward_map.get(u.pk),
        }
        rows.append(row)
    rows.sort(key=lambda r: (-r["score"], -r["days"], r["name"]))

    if request.method == "POST":
        try:
            uid = int(request.POST.get("user_id", ""))
        except ValueError:
            return HttpResponse("Invalid user.", status=400)
        row = next((r for r in rows if r["user"].pk == uid), None)
        if row is None:
            return HttpResponse("User not in selected report.", status=400)
        action = request.POST.get("action")
        with transaction.atomic():
            if action == "approve":
                gift = request.POST.get("gift", "").strip()
                if not row["enabled"] or row["score"] <= 0 or not gift or len(gift) > 255:
                    return HttpResponse("Positive score and reward description required.", status=400)
                reward, created = StaffPerformanceReward.objects.get_or_create(
                    user_id=uid, period_start=start, period_end=end,
                    defaults={
                        "score": row["score"],
                        "confirmed_contacts": row["contacts"],
                        "admissions": row["admissions"],
                        "gift_description": gift,
                        "approved_by": request.user,
                    },
                )
                messages.success(request, "Reward approved." if created else "Reward already recorded.")
            elif action == "given":
                reward = StaffPerformanceReward.objects.select_for_update().filter(
                    user_id=uid, period_start=start, period_end=end,
                ).first()
                if reward is None:
                    return HttpResponse("Approve reward first.", status=400)
                if reward.status == "approved":
                    reward.status = "given"
                    reward.fulfilled_at = timezone.now()
                    reward.fulfilled_by = request.user
                    reward.save(update_fields=["status", "fulfilled_at", "fulfilled_by"])
                messages.success(request, "Reward marked as given.")
            else:
                return HttpResponse("Invalid action.", status=400)
        return redirect(request.get_full_path())


    attendance_history = StaffAttendance.objects.filter(
        staff__user_id__in=ids,
        attendance_date__gte=start, attendance_date__lte=end,
    ).select_related("staff", "staff__user").order_by(
        "-attendance_date", "staff__user__username"
    )
    export = request.GET.get("export")
    headings = [
        "User", "Branch", "Role", "Work days", "No recorded work days",
        "Unique call leads", "Unique WhatsApp leads", "Confirmed contacts",
        "Results", "Admissions", "Receipts", "Collection",
        "Assigned today", "Assigned overdue", "Score",
    ]
    data = [
        [r["name"], r["branch"], r["role"], r["days"], r["quiet_days"],
         r["calls"], r["whatsapp"], r["contacts"], r["results"],
         r["admissions"], r["receipts"], float(r["collection"]),
         r["today"], r["overdue"], r["score"]]
        for r in rows
    ]

    headings.extend([
        "Attendance today", "Present days", "Leave days",
        "Recorded absent days", "Days without attendance record",
    ])
    for values, row in zip(data, rows):
        values.extend([
            row["attendance_status"], row["present_days"], row["leave_days"],
            row["absent_days"], row["missing_days"],
        ])
    if export == "excel":
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        wb = Workbook()
        ws = wb.active
        ws.title = "Team Monitoring"
        ws.append(headings)
        for values in data:
            ws.append(values)
            for cell in ws[ws.max_row]:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="171717")
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for column in ws.columns:
            ws.column_dimensions[column[0].column_letter].width = 23

        attendance_sheet = wb.create_sheet("Attendance Records")
        attendance_sheet.append([
            "Date", "User", "Username", "Branch", "Status",
            "First login", "Source",
        ])
        for entry in attendance_history:
            attendance_sheet.append([
                str(entry.attendance_date),
                entry.staff.user.get_full_name() or entry.staff.user.username,
                entry.staff.user.username,
                entry.branch,
                entry.get_status_display(),
                timezone.localtime(entry.first_login_time).strftime("%Y-%m-%d %H:%M")
                if entry.first_login_time else "",
                entry.source,
            ])
            for cell in attendance_sheet[attendance_sheet.max_row]:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
        attendance_sheet.freeze_panes = "A2"
        attendance_sheet.auto_filter.ref = attendance_sheet.dimensions
        for column in attendance_sheet.columns:
            attendance_sheet.column_dimensions[column[0].column_letter].width = 24
        buffer = BytesIO()
        wb.save(buffer)
        response = HttpResponse(
            buffer.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        response["Content-Disposition"] = 'attachment; filename="team-monitoring.xlsx"'
    elif export == "pdf":
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, LongTable, TableStyle
        buffer = BytesIO()
        styles = getSampleStyleSheet()
        small = styles["BodyText"].clone("small")
        small.fontSize = 8
        small.leading = 10
        table_data = [["User / Branch", "Work days", "Calls", "WhatsApp",
                       "Contacts", "Admissions", "Today", "Overdue", "Score"]]
        for r in rows:
            table_data.append([
                Paragraph(escape(r["name"]+" / "+r["branch"]), small),
                r["days"], r["calls"], r["whatsapp"], r["contacts"],
                r["admissions"], r["today"], r["overdue"], r["score"],
            ])
        table = LongTable(table_data, repeatRows=1,
                          colWidths=[210, 65, 55, 65, 65, 75, 55, 60, 55])
        table.setStyle(TableStyle([
            ("BACKGROUND", (0,0), (-1,0), colors.HexColor("#171717")),
            ("TEXTCOLOR", (0,0), (-1,0), colors.white),
            ("FONTNAME", (0,0), (-1,0), "Helvetica-Bold"),
            ("FONTSIZE", (0,0), (-1,-1), 8),
            ("VALIGN", (0,0), (-1,-1), "TOP"),
            ("BOTTOMPADDING", (0,0), (-1,-1), 8),
            ("TOPPADDING", (0,0), (-1,-1), 8),
            ("ROWBACKGROUNDS", (0,1), (-1,-1), [colors.white, colors.HexColor("#fff5eb")]),
            ("LINEBELOW", (0,0), (-1,0), 1, colors.HexColor("#ff7200")),
        ]))

        attendance_data = [[
            "User", "Today", "Present", "Leave", "Recorded absent", "No record",
        ]]
        for row in rows:
            attendance_data.append([
                Paragraph(escape(row["name"]), small),
                row["attendance_status"], row["present_days"], row["leave_days"],
                row["absent_days"],
                row["missing_days"] if row["missing_days"] is not None else "N/A",
            ])
        attendance_table = LongTable(
            attendance_data, repeatRows=1,
            colWidths=[220, 130, 90, 90, 100, 95],
        )
        attendance_table.setStyle(TableStyle([
            ("BACKGROUND", (0,0), (-1,0), colors.HexColor("#fff3e8")),
            ("FONTSIZE", (0,0), (-1,-1), 8),
            ("VALIGN", (0,0), (-1,-1), "TOP"),
            ("TOPPADDING", (0,0), (-1,-1), 8),
            ("BOTTOMPADDING", (0,0), (-1,-1), 8),
            ("LINEBELOW", (0,0), (-1,-1), .3, colors.HexColor("#e2e8f0")),
        ]))
        SimpleDocTemplate(buffer, pagesize=landscape(A4),
                          leftMargin=20, rightMargin=20).build([
            Paragraph("MCTI — Team Monitoring & Rewards", styles["Title"]),
            Paragraph(f"{start} to {end} | Score: confirmed contacts × 2 + admissions × 10", small),
            Spacer(1, 14), table, Spacer(1, 20),
            Paragraph("Staff Attendance", styles["Heading2"]),
            attendance_table, Spacer(1, 14),
            Paragraph("No recorded work is not absence. Calls/messages are initiated actions. "
                      "An Initiative by Maharashtra Computer Training Institute", small),
        ])
        response = HttpResponse(buffer.getvalue(), content_type="application/pdf")
        response["Content-Disposition"] = 'attachment; filename="team-monitoring.pdf"'
    else:
        params = request.GET.copy()
        params.pop("export", None)
        query = params.urlencode()
        recent = activities.select_related("created_by", "enquiry").order_by("-created_at")[:50]
        history = StaffPerformanceReward.objects.filter(
            user_id__in=ids
        ).select_related("user", "approved_by", "fulfilled_by")[:30]
        response = render(request, "core/team_monitoring.html", {
            **award_context,
            "rows": rows, "period": period, "start": start, "end": end,
            "branch": branch, "branches": branch_choices,
            "excel_query": query + ("&" if query else "") + "export=excel",
            "pdf_query": query + ("&" if query else "") + "export=pdf",
            "active": sum(r["days"] > 0 for r in rows),
            "quiet": sum(r["days"] == 0 for r in rows),
            "total": len(rows),
            "due_today": pending.filter(followup_date=today).count(),
            "due_overdue": pending.filter(followup_date__lt=today).count(),
            "unassigned": pending.filter(
                assigned_user__isnull=True, followup_date__lte=today
            ).count(),
            "recent": recent, "history": history,
            "attendance_history": attendance_history[:100],
            "attendance_date": today,
            "attendance_present": sum(
                r["attendance_status"] == "Present" for r in rows
            ),
        })
    response["Cache-Control"] = "private, no-store"
    return response
